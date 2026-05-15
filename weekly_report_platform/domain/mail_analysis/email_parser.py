"""邮件抓取与内容整理。

这个文件负责三件事：
1. 连接 POP3 邮箱并筛选候选邮件。
2. 解析邮件头、正文、HTML 和附件。
3. 把一封邮件整理成适合后续 LLM 分析的统一文本。
"""

from __future__ import annotations

import email
import io
import os
import poplib
import re
import tempfile
import warnings
from datetime import datetime
from email.header import decode_header
from email.utils import parsedate_to_datetime
from typing import Any

from markitdown import MarkItDown
from openpyxl import load_workbook

from weekly_report_platform.infrastructure.config import MailSettings, load_project_env
from weekly_report_platform.infrastructure.logging import logger


class EmailService:
    """邮箱访问服务。

    这个类封装了 POP3 登录、邮件筛选和单封邮件解析，
    让上层不需要直接接触 `poplib` 和 MIME 细节。
    """

    def __init__(self) -> None:
        # 初始化时只加载配置，不立即联网。
        # 这样测试时更容易替换行为，也避免导入模块就触发外部连接。
        load_project_env()
        settings = MailSettings()
        self.pop3_host = settings.pop3_host
        self.pop3_port = settings.pop3_port
        self.username = settings.username
        self.password = settings.password
        self.subject_keyword = settings.subject_keyword
        self.pop_conn = None

    def login_pop3(self) -> bool:
        """登录 POP3 邮箱。成功返回 True，失败返回 False。"""
        try:
            self.pop_conn = poplib.POP3_SSL(self.pop3_host, self.pop3_port, timeout=30)
            self.pop_conn.user(self.username)
            self.pop_conn.pass_(self.password)
            logger.info(f"Logged into POP3 server: {self.pop3_host}")
            return True
        except Exception as exc:
            logger.error(f"Failed to login to POP3 server: {exc}")
            return False

    def logout(self) -> None:
        """退出邮箱连接。"""
        if self.pop_conn:
            try:
                self.pop_conn.quit()
                logger.info("Logged out of POP3 server")
            except Exception as exc:
                logger.error(f"Failed to logout from POP3 server: {exc}")
            finally:
                self.pop_conn = None

    def get_email_count(self) -> int:
        """返回邮箱中的邮件总数。"""
        if not self.pop_conn:
            logger.error("POP3 connection is not initialized")
            return 0
        try:
            return len(self.pop_conn.list()[1])
        except Exception as exc:
            logger.error(f"Failed to get email count: {exc}")
            return 0

    def _decode_header_bytes(self, data: bytes, declared_encoding: str | None) -> str:
        """把邮件头中的 bytes 片段解码成字符串。"""
        candidate_encodings = [declared_encoding, "utf-8", "gb18030", "gbk", "gb2312"]
        for encoding in candidate_encodings:
            if not encoding:
                continue
            try:
                return data.decode(encoding)
            except Exception:
                continue
        return data.decode("utf-8", errors="replace")

    def _decode_header_value(self, header_value) -> str:
        """解码邮件头字段。

        某些邮件头会由多个不同编码片段拼接而成，
        这里统一把它们还原成一条可读字符串。
        """
        if not header_value:
            return ""
        decoded_parts = decode_header(header_value)
        result_parts = []
        for data, encoding in decoded_parts:
            if isinstance(data, bytes):
                result_parts.append(self._decode_header_bytes(data, encoding))
            else:
                result_parts.append(data)
        return "".join(result_parts)

    def _extract_message_id_from_msg(self, message) -> str:
        """提取并清理 Message-ID。"""
        mid = message.get("Message-ID", "")
        if not mid:
            return ""
        mid_str = str(mid).strip().replace("\r", "").replace("\n", "")
        return re.sub(r"\s+", " ", mid_str)

    def get_email_header_only(self, message_id: int) -> dict[str, str] | None:
        """只拉取邮件头。

        扫描邮箱时先只看头信息，可以显著减少网络和解析开销。
        """
        if not self.pop_conn:
            logger.error("POP3 connection is not initialized")
            return None
        try:
            _, lines, _ = self.pop_conn.top(message_id, 0)
            raw_headers = b"\r\n".join(lines)
            message = email.message_from_bytes(raw_headers)
            return {
                "subject": self._decode_header_value(message.get("Subject", "")),
                "date": message.get("Date", ""),
                "message_id": self._extract_message_id_from_msg(message),
            }
        except Exception as exc:
            logger.debug(f"Failed to get email header {message_id}: {exc}")
            return None

    def _decode_payload(self, payload_bytes: bytes, charset: str) -> str:
        """解码正文内容。"""
        for encoding in [charset, "gb2312", "gbk", "gb18030", "utf-8"]:
            try:
                return payload_bytes.decode(encoding)
            except Exception:
                continue
        return payload_bytes.decode("utf-8", errors="replace")

    def get_email_by_id(self, message_id: int) -> dict[str, Any] | None:
        """拉取并解析一封完整邮件。"""
        if not self.pop_conn:
            logger.error("POP3 connection is not initialized")
            return None
        try:
            _, lines, _ = self.pop_conn.retr(message_id)
            raw_email = b"\r\n".join(lines)
            email_data = self._parse_email(raw_email)
            if email_data:
                return {"id": str(message_id), **email_data}
            return None
        except Exception as exc:
            logger.error(f"Failed to get email {message_id}: {exc}")
            return None

    def _parse_email(self, raw_email_bytes: bytes) -> dict[str, Any] | None:
        """解析原始 MIME 邮件。

        返回一个 dict，包含：
        - `subject` / `from` / `to` / `cc` / `date`
        - `text_body` / `html_body`
        - `message_id`
        - `attachments`
        """
        try:
            message = email.message_from_bytes(raw_email_bytes)
            subject = self._decode_header_value(message.get("Subject", ""))
            from_addr = self._decode_header_value(message.get("From", ""))
            to_addr = self._decode_header_value(message.get("To", ""))
            cc_addr = self._decode_header_value(message.get("Cc", ""))
            message_id = self._extract_message_id_from_msg(message)
            date = message.get("Date", "")

            text_body = ""
            html_body = ""
            attachments = []
            for part in message.walk():
                # 按“纯文本正文 / HTML 正文 / 附件”三类处理。
                content_type = part.get_content_type()
                content_disposition = str(part.get("Content-Disposition", ""))

                if content_type == "text/plain" and "attachment" not in content_disposition:
                    charset = part.get_content_charset() or "utf-8"
                    payload_bytes = part.get_payload(decode=True)
                    if payload_bytes:
                        text_body += self._decode_payload(payload_bytes, charset)
                elif content_type == "text/html" and "attachment" not in content_disposition:
                    charset = part.get_content_charset() or "utf-8"
                    payload_bytes = part.get_payload(decode=True)
                    if payload_bytes:
                        html_body += self._decode_payload(payload_bytes, charset)
                elif content_type == "application/octet-stream" or "attachment" in content_disposition:
                    filename = part.get_filename()
                    if filename:
                        filename = self._decode_header_value(filename)
                        decoded_data = part.get_payload(decode=True)
                        if decoded_data:
                            attachments.append(
                                {
                                    "filename": filename,
                                    "content_type": content_type,
                                    "data": decoded_data,
                                    "size": len(decoded_data),
                                }
                            )

            return {
                "subject": subject,
                "from": from_addr,
                "to": to_addr,
                "cc": cc_addr,
                "date": date,
                "text_body": text_body,
                "html_body": html_body,
                "message_id": message_id,
                "attachments": attachments,
            }
        except Exception as exc:
            logger.error(f"Failed to parse email: {exc}")
            return None

    def _parse_date(self, date_str: str) -> datetime | None:
        """解析邮件头里的 Date 字段。"""
        if not date_str:
            return None
        try:
            return parsedate_to_datetime(date_str)
        except Exception as exc:
            logger.debug(f"Failed to parse date {date_str}: {exc}")
            return None

    def find_emails_by_filter(
        self,
        start_date: datetime | None,
        subject_keyword: str | None = None,
        max_emails: int | None = None,
    ) -> list[dict[str, Any]]:
        """按起始日期和主题关键字筛选邮件。"""
        if not self.pop_conn:
            logger.error("POP3 connection is not initialized")
            return []

        keyword = subject_keyword or self.subject_keyword
        total = self.get_email_count()
        logger.info(f"Scanning mailbox: total={total}, start_date={start_date}, keyword={keyword}, max_emails={max_emails}")

        matched_emails = []
        stop_threshold = 10
        old_email_count = 0

        for email_id in range(total, 0, -1):
            # 从新到旧扫描。
            # 如果连续遇到很多封早于起始日期的邮件，就提前停止，
            # 避免每次都把整个大邮箱扫穿。
            if old_email_count >= stop_threshold:
                break
            header_info = self.get_email_header_only(email_id)
            if not header_info:
                continue

            email_date = self._parse_date(header_info.get("date", ""))
            if start_date is not None and email_date:
                try:
                    if email_date < start_date:
                        old_email_count += 1
                        continue
                except TypeError:
                    if email_date.replace(tzinfo=None) < start_date.replace(tzinfo=None):
                        old_email_count += 1
                        continue

            old_email_count = 0
            subject = header_info.get("subject", "")
            if keyword and keyword in subject:
                email_data = self.get_email_by_id(email_id)
                if email_data:
                    matched_emails.append(email_data)
                    if max_emails is not None and len(matched_emails) >= max_emails:
                        break

        logger.info(f"Matched emails: {len(matched_emails)}")
        return matched_emails


def format_email_for_analysis(email_data: dict[str, Any]) -> str:
    """把解析后的邮件对象转换成给大模型使用的统一文本。

    输入 `email_data` 通常包含：
    - `subject` / `from` / `to` / `cc` / `date`
    - `text_body` / `html_body`
    - `attachments`

    返回值：
    - 一个 `str`
    - 结构上会分成邮件元信息、正文、附件三大段
    - 如果存在附件，会把每个可解析附件的文本都拼接进去
    """
    content_parts = []

    content_parts.append("=" * 60)
    content_parts.append("[Email Metadata]")
    content_parts.append("=" * 60)
    content_parts.append(f"Subject: {email_data.get('subject', 'Unknown')}")
    content_parts.append(f"From: {email_data.get('from', 'Unknown')}")
    content_parts.append(f"Date: {email_data.get('date', 'Unknown')}")
    content_parts.append("")

    content_parts.append("=" * 60)
    content_parts.append("[Email Body]")
    content_parts.append("=" * 60)

    html_body = email_data.get("html_body", "").strip()
    text_body = email_data.get("text_body", "").strip()
    if html_body:
        # 优先使用 HTML，因为很多周报格式化信息只存在于 HTML 正文里。
        markdown_body = _html_to_markdown(html_body)
        if markdown_body:
            content_parts.append(markdown_body)
        else:
            clean_text = re.sub(r"<[^>]+>", " ", html_body)
            content_parts.append(re.sub(r"\s+", " ", clean_text).strip())
    elif text_body:
        content_parts.append(text_body)
    else:
        content_parts.append("[Empty email body; this email may contain attachments only]")
    content_parts.append("")

    attachments = email_data.get("attachments", [])
    if attachments:
        for attachment in attachments:
            filename = attachment.get("filename", "attachment")
            extension = os.path.splitext(filename)[1].lower()
            content_parts.append("=" * 60)
            content_parts.append(f"[Attachment: {filename}]")
            content_parts.append("=" * 60)
            if extension in [".xlsx", ".xls"]:
                content_parts.append(_parse_excel_attachment(attachment) or "[Failed to parse Excel attachment]")
            elif extension in [".docx", ".doc"]:
                content_parts.append(_parse_word_attachment(attachment) or "[Failed to parse Word attachment]")
            elif extension == ".pdf":
                content_parts.append(_parse_pdf_attachment(attachment) or "[Failed to parse PDF attachment]")
            elif extension == ".txt":
                content_parts.append(_parse_text_attachment(attachment) or "[Empty text attachment]")
            else:
                # 不支持的附件类型只标记出来，不让整封邮件失败。
                content_parts.append(f"[Unsupported attachment type: {extension}]")
            content_parts.append("")
    else:
        content_parts.append("=" * 60)
        content_parts.append("[Attachments]")
        content_parts.append("=" * 60)
        content_parts.append("[No attachments]")
        content_parts.append("")

    return "\n".join(content_parts)


def _parse_excel_attachment(attachment: dict[str, Any]) -> str:
    """解析 Excel 附件。

    当前策略是把所有 sheet 都转成 Markdown 表格并拼起来，
    这样后续模型可以看到整份表格内容。
    """
    try:
        data = attachment.get("data")
        if not data:
            return ""
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
            workbook = load_workbook(io.BytesIO(data), data_only=True)
        sections: list[str] = []
        for sheet in workbook.worksheets:
            rows = [[_process_cell_value(cell) for cell in row] for row in sheet.iter_rows()]
            markdown_table = _rows_to_markdown_table(rows)
            if not markdown_table:
                continue
            sections.append(f"### Sheet: {sheet.title}\n\n{markdown_table}")
        workbook.close()
        return "\n\n".join(sections)
    except Exception as exc:
        logger.error(f"Failed to parse Excel attachment: {exc}")
        return ""


def _parse_word_attachment(attachment: dict[str, Any]) -> str:
    """解析 Word 附件。"""
    from docx import Document

    try:
        data = attachment.get("data")
        if not data:
            return ""
        doc = Document(io.BytesIO(data))
        return "\n".join([paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()])
    except Exception as exc:
        logger.error(f"Failed to parse Word attachment: {exc}")
        return ""


def _parse_pdf_attachment(attachment: dict[str, Any]) -> str:
    """解析 PDF 附件。"""
    import fitz

    try:
        data = attachment.get("data")
        if not data:
            return ""
        doc = fitz.open(stream=data, filetype="pdf")
        text_content = []
        for index in range(len(doc)):
            text_content.append(f"\n=== Page {index + 1} ===\n")
            text_content.append(doc[index].get_text())
        doc.close()
        return "".join(text_content)
    except Exception as exc:
        logger.error(f"Failed to parse PDF attachment: {exc}")
        return ""


def _parse_text_attachment(attachment: dict[str, Any]) -> str:
    """解析纯文本附件。"""
    try:
        data = attachment.get("data")
        if not data:
            return ""
        for encoding in ["utf-8", "gb18030", "gbk", "gb2312", "latin-1"]:
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="ignore")
    except Exception as exc:
        logger.error(f"Failed to parse text attachment: {exc}")
        return ""


def _process_cell_value(cell) -> str:
    """把 Excel 单元格值标准化成 Markdown 友好的文本。"""
    if cell.value is None:
        return ""
    cell_format = cell.number_format
    if "%" in str(cell_format) and isinstance(cell.value, (int, float)):
        result = f"{cell.value * 100:.2f}%"
    else:
        result = str(cell.value)
    result = result.replace("|", "\\|").replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")
    return result.strip()


def _rows_to_markdown_table(rows: list[list[str]]) -> str:
    """把二维单元格数组转成 Markdown 表格。"""
    if not rows:
        return ""
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        return ""
    num_cols = len(rows[0])
    md_lines = []
    for index, row in enumerate(rows):
        while len(row) < num_cols:
            row.append("")
        md_lines.append("|" + "|".join(row[:num_cols]) + "|")
        if index == 0:
            md_lines.append("|" + "|".join(["---"] * num_cols) + "|")
    return "\n".join(md_lines)


def _html_to_markdown(html_content: str) -> str:
    """把 HTML 正文转成 Markdown。"""
    temp_html_path = None
    try:
        html_content = html_content.replace("charset=gb2312", "charset=utf-8").replace("charset=GB2312", "charset=utf-8")
        html_content = html_content.replace("charset=gbk", "charset=utf-8").replace("charset=GBK", "charset=utf-8")
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", suffix=".html", delete=False) as file:
            temp_html_path = file.name
            file.write(html_content)
        md = MarkItDown(enable_plugins=False)
        result = md.convert(temp_html_path)
        return _clean_markdown_content(result.text_content)
    except Exception as exc:
        logger.error(f"MarkItDown conversion failed: {exc}")
        return ""
    finally:
        if temp_html_path:
            try:
                os.remove(temp_html_path)
            except Exception:
                pass


def _clean_markdown_content(markdown_text: str) -> str:
    """清洗 Markdown 文本，减少空白和 HTML 实体噪声。"""
    if not markdown_text:
        return ""
    markdown_text = re.sub(r"(&nbsp;)+", " ", markdown_text)
    markdown_text = re.sub(r"\u00a0+", " ", markdown_text)
    markdown_text = markdown_text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"')
    lines = [line.rstrip() for line in markdown_text.split("\n")]
    cleaned_lines = []
    empty_count = 0
    for line in lines:
        if line == "":
            empty_count += 1
            if empty_count <= 2:
                cleaned_lines.append(line)
        else:
            empty_count = 0
            cleaned_lines.append(line)
    result = "\n".join(cleaned_lines).strip()
    return re.sub(r" {3,}", "  ", result)


email_service = EmailService()
